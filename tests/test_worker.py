from __future__ import annotations

import base64
import json
from pathlib import Path

import pytest

from tests.helpers import (
    OVERCLAIM_CLAIM_QUOTE,
    OVERCLAIM_EVIDENCE_QUOTE,
    sample_history_v1,
    sample_long_section_draft,
    sample_new_draft,
    sample_overclaim_draft,
    sample_same_heading_fixed_body,
)
from thesis_review.errors import ReviewError
from thesis_review.history.store import HistoryStore
from thesis_review.service import ThesisReviewService
from thesis_review.worker import NAV_BUDGET, Worker
from thesis_review.word.adapter import WordAdapter


def _open(worker: Worker, data: bytes) -> None:
    worker.dispatch("open_draft", {"bytes_b64": base64.b64encode(data).decode("ascii")})


def _paragraph_containing(worker: Worker, needle: str) -> dict:
    for item in worker.dispatch("list_paragraphs", {})["paragraphs"]:
        if needle in item["text"]:
            return item
    raise AssertionError(f"missing paragraph containing {needle!r}")


def test_worker_lists_paragraphs_from_open_draft(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    payload = base64.b64encode(sample_new_draft()).decode("ascii")
    opened = worker.dispatch("open_draft", {"bytes_b64": payload})
    assert opened["n_paragraphs"] >= 3
    paragraphs = worker.dispatch("list_paragraphs", {})["paragraphs"]
    assert any("非常非常有效" in item["text"] for item in paragraphs)


def test_worker_live_events_exist_before_commit_without_claim_quotes(tmp_path: Path):
    live = tmp_path / "out" / "live"
    worker = Worker(
        home=tmp_path,
        teacher_id="teacher-a",
        student_id="zhou",
        major="人工智能",
        live=live,
    )
    _open(worker, sample_overclaim_draft())
    worker.dispatch("list_outline", {})
    heading = _paragraph_containing(worker, "3 实验结果")
    worker.dispatch("read_section", {"start_ordinal": heading["ordinal"], "limit": 4})
    worker.dispatch(
        "record_argument_finding",
        {
            "claim_quote": OVERCLAIM_CLAIM_QUOTE,
            "evidence_quote": OVERCLAIM_EVIDENCE_QUOTE,
            "problem": "结论用词过满，缺少显著性检验支持。",
            "rationale": "提升幅度与用词不符。",
            "draft_id": "overclaim",
        },
    )
    events_path = live / "events.jsonl"
    assert events_path.is_file()
    events = [json.loads(line) for line in events_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    ops = [item["op"] for item in events]
    assert "open_draft" in ops
    assert "list_outline" in ops
    assert "read_section" in ops
    assert "record_argument_finding" in ops
    assert "commit_review" not in ops
    dumped = json.dumps(events, ensure_ascii=False)
    assert OVERCLAIM_CLAIM_QUOTE not in dumped
    assert OVERCLAIM_EVIDENCE_QUOTE not in dumped
    assert "api_key" not in dumped
    read_event = next(item for item in events if item["op"] == "read_section")
    assert "3 实验结果" in str(read_event.get("heading") or "")
    snapshot = json.loads((live / "findings.json").read_text(encoding="utf-8"))
    assert snapshot
    assert any(item.get("source") == "argument" for item in snapshot)


def test_worker_commit_review_writes_findings(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    worker.dispatch("open_draft", {"bytes_b64": base64.b64encode(sample_new_draft()).decode("ascii")})
    worker.dispatch("run_checks", {"draft_id": "new"})
    result = worker.dispatch("commit_review", {"draft_id": "new", "output_dir": str(tmp_path / "out")})
    assert Path(result["reviewed_path"]).is_file()
    assert Path(result["findings_path"]).is_file()


def test_commit_review_writes_trace_without_quotes_or_body(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    worker.dispatch("list_outline", {})
    heading = _paragraph_containing(worker, "3 实验结果")
    worker.dispatch("read_section", {"start_ordinal": heading["ordinal"], "limit": 4})
    worker.dispatch(
        "record_argument_finding",
        {
            "claim_quote": OVERCLAIM_CLAIM_QUOTE,
            "evidence_quote": OVERCLAIM_EVIDENCE_QUOTE,
            "problem": "结论用词过满，缺少显著性检验支持。",
            "rationale": "提升幅度与用词不符。",
            "draft_id": "overclaim",
        },
    )
    result = worker.dispatch(
        "commit_review",
        {"draft_id": "overclaim", "output_dir": str(tmp_path / "out")},
    )
    trace_path = Path(result["trace_path"])
    assert trace_path.is_file()
    payload = json.loads(trace_path.read_text(encoding="utf-8"))
    ops = [item["op"] for item in payload["ops"]]
    assert "list_outline" in ops
    assert "read_section" in ops
    assert "record_argument_finding" in ops
    assert "commit_review" in ops
    dumped = json.dumps(payload, ensure_ascii=False)
    assert OVERCLAIM_CLAIM_QUOTE not in dumped
    assert OVERCLAIM_EVIDENCE_QUOTE not in dumped
    assert "api_key" not in dumped
    read_entry = next(item for item in payload["ops"] if item["op"] == "read_section")
    assert read_entry["ok"] is True
    assert read_entry["params"]["start_ordinal"] == heading["ordinal"]
    record_entry = next(item for item in payload["ops"] if item["op"] == "record_argument_finding")
    assert "claim_quote" not in record_entry.get("params", {})
    assert "evidence_quote" not in record_entry.get("params", {})


def test_trace_records_failed_nav_and_open_draft_resets(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    for _unused in range(worker.nav_budget):
        worker.dispatch("list_outline", {})
    with pytest.raises(ReviewError):
        worker.dispatch("list_outline", {})
    failed = [item for item in worker.trace if item["op"] == "list_outline" and item["ok"] is False]
    assert failed
    assert failed[-1]["code"] == "nav_budget"
    _open(worker, sample_new_draft())
    assert worker.trace
    assert worker.trace[-1]["op"] == "open_draft"
    assert all(item["op"] == "open_draft" for item in worker.trace)


def _seed_confirmed_issue(home: Path, *, data: bytes, needle: str) -> str:
    service = ThesisReviewService(
        store=HistoryStore(home / "thesis-review.sqlite"),
        adapter=WordAdapter(),
        home=home,
    )
    candidates = service.ingest_history(
        teacher_id="teacher-a",
        student_id="zhou",
        major="人工智能",
        draft_id="v1",
        data=data,
    )
    issue = next(item for item in candidates if needle in item.original_text)
    service.confirm_issue(teacher_id="teacher-a", student_id="zhou", issue_id=issue.id)
    return issue.id


def test_get_history_candidates_recalls_without_writing_findings(tmp_path: Path):
    issue_id = _seed_confirmed_issue(tmp_path, data=sample_history_v1(), needle="主观评价")
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_new_draft())
    result = worker.dispatch("get_history_candidates", {"draft_id": "new"})
    candidates = result["candidates"]
    assert candidates
    hit = next(item for item in candidates if item["issue_id"] == issue_id)
    assert hit["category"]
    assert "非常非常有效" in hit["new_quote"]
    assert hit["new_anchor"].startswith("P")
    assert "避免主观评价" in (hit["original_text"] or hit["problem"])
    assert hit["old_span"]
    assert worker.findings == []
    assert WordAdapter().extract_comments(worker.opened) == []


def test_confirm_history_finding_rejects_invented_quote(tmp_path: Path):
    issue_id = _seed_confirmed_issue(tmp_path, data=sample_history_v1(), needle="主观评价")
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_new_draft())
    with pytest.raises(ReviewError) as missing:
        worker.dispatch(
            "confirm_history_finding",
            {"issue_id": issue_id, "new_quote": "", "rationale": "仍是主观评价。"},
        )
    assert missing.value.code == "missing_quote"
    with pytest.raises(ReviewError) as invented:
        worker.dispatch(
            "confirm_history_finding",
            {
                "issue_id": issue_id,
                "new_quote": "本节仍缺少评价指标。",
                "rationale": "编造的原文。",
            },
        )
    assert invented.value.code == "quote_not_in_draft"
    assert worker.findings == []
    assert all(item.source != "history" for item in worker.findings)


def test_confirm_history_finding_accepts_real_quote(tmp_path: Path):
    issue_id = _seed_confirmed_issue(tmp_path, data=sample_history_v1(), needle="主观评价")
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_new_draft())
    candidates = worker.dispatch("get_history_candidates", {"draft_id": "new"})["candidates"]
    hit = next(item for item in candidates if item["issue_id"] == issue_id)
    result = worker.dispatch(
        "confirm_history_finding",
        {
            "issue_id": issue_id,
            "new_quote": hit["new_quote"],
            "rationale": "仍缺实验依据。",
            "draft_id": "new",
        },
    )
    assert result["ok"] is True
    finding = worker.findings[-1]
    assert finding.source == "history"
    assert finding.issue_id == issue_id
    assert finding.teacher_decision == "pending"
    assert "仍出现已确认的历史问题" in finding.problem
    committed = worker.dispatch("commit_review", {"draft_id": "new", "output_dir": str(tmp_path / "out")})
    comments = WordAdapter().extract_comments(WordAdapter().open_path(committed["reviewed_path"]))
    blob = "\n".join(item.text for item in comments)
    assert "历次" not in blob
    assert hit["new_quote"] not in blob


def test_confirm_history_finding_accepts_source_quote_after_rule_revision(tmp_path: Path):
    issue_id = _seed_confirmed_issue(tmp_path, data=sample_history_v1(), needle="主观评价")
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_new_draft())
    worker.dispatch("run_checks", {"draft_id": "new"})
    candidates = worker.dispatch("get_history_candidates", {"draft_id": "new"})["candidates"]
    hit = next(item for item in candidates if item["issue_id"] == issue_id)
    result = worker.dispatch(
        "confirm_history_finding",
        {
            "issue_id": issue_id,
            "new_quote": hit["new_quote"],
            "rationale": "规则修订后仍是同一处主观评价。",
            "draft_id": "new",
        },
    )
    assert result["ok"] is True
    assert any(item.source == "history" for item in worker.findings)


def test_get_history_candidates_keeps_heading_hit_without_writing(tmp_path: Path):
    _seed_confirmed_issue(
        tmp_path, data=sample_same_heading_fixed_body(history=True), needle="实验步骤"
    )
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_same_heading_fixed_body(history=False))
    result = worker.dispatch("get_history_candidates", {"draft_id": "new"})
    assert result["candidates"]
    assert any("3.1 实验设计" in item["new_quote"] for item in result["candidates"])
    assert worker.findings == []


def test_list_outline_returns_headings_not_body(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    outline = worker.dispatch("list_outline", {})["outline"]
    texts = [item["text"] for item in outline]
    assert any("3 实验结果" in text for text in texts)
    assert any("4 结论" in text for text in texts)
    assert all(OVERCLAIM_CLAIM_QUOTE not in text for text in texts)
    assert all(OVERCLAIM_EVIDENCE_QUOTE not in text for text in texts)


def test_read_paragraphs_and_read_section_are_bounded(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_long_section_draft())
    heading = _paragraph_containing(worker, "1 实验")
    body = _paragraph_containing(worker, "实验步骤说明段落00")
    paragraphs = worker.dispatch(
        "read_paragraphs",
        {"start_ordinal": body["ordinal"], "limit": 20},
    )["paragraphs"]
    assert 1 <= len(paragraphs) <= 8
    section = worker.dispatch("read_section", {"start_ordinal": heading["ordinal"], "limit": 20})
    assert 1 <= len(section["paragraphs"]) <= 8
    assert section.get("truncated") is True


def test_find_text_returns_limited_hits_with_context(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    hits = worker.dispatch("find_text", {"needle": "0.81"})["hits"]
    assert hits
    assert len(hits) <= 5
    assert any("0.81" in item["snippet"] for item in hits)
    assert all("anchor" in item and "ordinal" in item for item in hits)


def test_record_argument_finding_rejects_missing_or_invented_quotes(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    with pytest.raises(ReviewError) as missing:
        worker.dispatch(
            "record_argument_finding",
            {
                "claim_quote": "",
                "evidence_quote": OVERCLAIM_EVIDENCE_QUOTE,
                "problem": "结论用词过满，缺少显著性检验支持。",
                "rationale": "缺少主张原文。",
            },
        )
    assert missing.value.code == "quote_not_in_draft"
    with pytest.raises(ReviewError) as invented:
        worker.dispatch(
            "record_argument_finding",
            {
                "claim_quote": OVERCLAIM_CLAIM_QUOTE,
                "evidence_quote": "准确率由 0.50 提高到 0.99，且差异极显著。",
                "problem": "结论用词过满，缺少显著性检验支持。",
                "rationale": "编造的结果句。",
            },
        )
    assert invented.value.code == "quote_not_in_draft"
    assert worker.findings == []


def test_record_argument_finding_accepts_real_quotes(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    result = worker.dispatch(
        "record_argument_finding",
        {
            "claim_quote": OVERCLAIM_CLAIM_QUOTE,
            "evidence_quote": OVERCLAIM_EVIDENCE_QUOTE,
            "problem": "结论用词过满，实验结果仅有微弱数值变化。",
            "rationale": "未见显著性检验，提升幅度与「显著提升」不符。",
            "draft_id": "new",
        },
    )
    assert result["ok"] is True
    finding = worker.findings[-1]
    assert finding.category == "B"
    assert finding.source == "argument"
    assert finding.code == "claim_without_evidence"
    assert finding.apply == "comment"
    assert finding.teacher_decision == "pending"
    committed = worker.dispatch("commit_review", {"draft_id": "new", "output_dir": str(tmp_path / "out")})
    comments = WordAdapter().extract_comments(WordAdapter().open_path(committed["reviewed_path"]))
    blob = "\n".join(item.text for item in comments)
    assert OVERCLAIM_CLAIM_QUOTE not in blob
    assert OVERCLAIM_EVIDENCE_QUOTE not in blob


def test_nav_budget_then_only_record_or_commit(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    for _unused in range(worker.nav_budget):
        worker.dispatch("list_outline", {})
    with pytest.raises(ReviewError) as caught:
        worker.dispatch("list_outline", {})
    assert caught.value.code == "nav_budget"
    assert str(worker.nav_budget) in str(caught.value)
    with pytest.raises(ReviewError):
        worker.dispatch("read_section", {"start_ordinal": 1})
    recorded = worker.dispatch(
        "record_argument_finding",
        {
            "claim_quote": OVERCLAIM_CLAIM_QUOTE,
            "evidence_quote": OVERCLAIM_EVIDENCE_QUOTE,
            "problem": "结论用词过满，缺少显著性检验支持。",
            "rationale": "导航次数用尽后仍可用已核对原文记录。",
        },
    )
    assert recorded["ok"] is True
    committed = worker.dispatch("commit_review", {"draft_id": "new", "output_dir": str(tmp_path / "out")})
    assert Path(committed["reviewed_path"]).is_file()


def test_worker_word_tools_are_blocked_by_teacher_gate(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_new_draft())
    with pytest.raises(ReviewError) as add:
        worker.dispatch("add_comment", {"anchor": "P1", "text": "直接写给学生"})
    assert add.value.code == "teacher_gate"
    with pytest.raises(ReviewError) as replace:
        worker.dispatch("replace_tracked", {"anchor": "P1", "old": "非常非常", "new": "较为"})
    assert replace.value.code == "teacher_gate"
    committed = worker.dispatch("commit_review", {"draft_id": "new", "output_dir": str(tmp_path / "out")})
    comments = WordAdapter().extract_comments(WordAdapter().open_path(committed["reviewed_path"]))
    assert comments == []


def test_worker_prefers_authoritative_draft_over_model_params(tmp_path: Path):
    draft_id = "基于LDA的热点舆情分析——以赤峰“免费菜事件”为例（第1稿）"
    out = tmp_path / "结果"
    source = out / f"{draft_id}-source.docx"
    out.mkdir(parents=True)
    source.write_bytes(sample_new_draft())
    worker = Worker(
        home=tmp_path,
        teacher_id="teacher-a",
        student_id="zhou",
        major="人工智能",
        draft_id=draft_id,
        draft_path=str(source),
        review_dir=out,
    )
    # The model normalizes curly quotes to '"' when relaying paths; the worker
    # must open the parent-provided file instead of the mangled param.
    opened = worker.dispatch("open_draft", {"path": str(out / '以赤峰"免费菜事件"为例.docx')})
    assert opened["n_paragraphs"] >= 3
    worker.dispatch("run_checks", {"draft_id": "mangled"})
    committed = worker.dispatch(
        "commit_review",
        {"draft_id": "以赤峰免费菜事件为例", "output_dir": str(tmp_path / "elsewhere")},
    )
    assert committed["findings_path"] == str(out / f"{draft_id}-findings.json")
    assert Path(committed["findings_path"]).is_file()
    assert not (tmp_path / "elsewhere").exists()
    findings = json.loads(Path(committed["findings_path"]).read_text(encoding="utf-8"))
    assert findings
    assert all(item["draft_id"] == draft_id for item in findings)


def test_nav_budget_scales_with_long_draft_and_reports_nav_left(tmp_path: Path):
    from docx import Document

    from tests.helpers import _save

    doc = Document()
    doc.add_paragraph("本科毕业论文")
    for index in range(340):
        doc.add_paragraph(f"第{index}段：本研究围绕热点舆情展开分析。")
    data = _save(doc)
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    opened = worker.dispatch("open_draft", {"bytes_b64": base64.b64encode(data).decode("ascii")})
    assert opened["n_paragraphs"] >= 340
    # 20 reads x 8 paragraphs cannot cover a 340-paragraph thesis; the budget
    # grows with the draft (43 here) instead of cutting the agent off mid-paper.
    assert worker.nav_budget == 43
    assert opened["nav_budget"] == 43
    first = worker.dispatch("list_outline", {})
    assert first["nav_left"] == 42
    second = worker.dispatch("read_paragraphs", {"start_ordinal": 1})
    assert second["nav_left"] == 41
    hits = worker.dispatch("find_text", {"needle": "热点舆情"})
    assert "nav_left" in hits
    assert hits["nav_left"] == 40


def test_record_finding_rejects_freeform_kind(tmp_path: Path):
    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    para = _paragraph_containing(worker, OVERCLAIM_CLAIM_QUOTE)
    with pytest.raises(ReviewError) as caught:
        worker.dispatch(
            "record_content_finding",
            {
                "kind": "语言问题",
                "subtype": "结论无支撑",
                "quote": para["text"][:12],
                "problem": "结论无支撑。",
                "rationale": "缺少实验依据。",
            },
        )
    assert caught.value.code == "invalid_params"
    assert "content/external/format/language" in str(caught.value)
    ok = worker.dispatch(
        "record_content_finding",
        {
            "kind": "language",
            "subtype": "terminology",
            "quote": para["text"][:12],
            "problem": "关键用词缺少实验定义。",
            "rationale": "缺少实验依据。",
        },
    )
    assert ok["ok"] is True
    assert worker.findings[-1].category == "A"


def test_argument_limit_message_names_external_at_cap(tmp_path: Path):
    from thesis_review.worker import MAX_CONTENT_FINDINGS

    worker = Worker(home=tmp_path, teacher_id="teacher-a", student_id="zhou", major="人工智能")
    _open(worker, sample_overclaim_draft())
    para = _paragraph_containing(worker, OVERCLAIM_CLAIM_QUOTE)
    quote = para["text"][:10]
    for _index in range(MAX_CONTENT_FINDINGS):
        worker.dispatch(
            "record_content_finding",
            {
                "kind": "content",
                "subtype": "structure",
                "quote": quote,
                "problem": "结论缺少可核验依据支撑。",
                "rationale": "需要实验数据支持。",
            },
        )
    with pytest.raises(ReviewError) as caught:
        worker.dispatch(
            "record_external_finding",
            {
                "quote": quote,
                "problem": "外部数据需核对。",
                "rationale": "宏观数据需要来源。",
                "external_sources": [{"title": "来源", "url": "https://example.com/a"}],
            },
        )
    assert caught.value.code == "argument_limit"
    assert "外部核验" in str(caught.value)
    assert str(MAX_CONTENT_FINDINGS) in str(caught.value)
