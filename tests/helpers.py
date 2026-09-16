"""Shared in-memory DOCX builders for tests. No real student papers."""
from __future__ import annotations

import io

from docx import Document

from thesis_review.fixtures import (
    FULL_ABSTRACT_RATE,
    FULL_CONCLUSION_RATE,
    FULL_OVERCLAIM,
    FULL_STATS,
    OVERCLAIM_CLAIM_QUOTE,
    OVERCLAIM_EVIDENCE_QUOTE,
    SUPPORTED_CLAIM_QUOTE,
    SUPPORTED_EVIDENCE_QUOTE,
    full_thesis_draft as sample_full_thesis_draft,
    overclaim_draft as sample_overclaim_draft,
    supported_claim_draft as sample_supported_claim_draft,
)

__all__ = [
    "OVERCLAIM_CLAIM_QUOTE",
    "OVERCLAIM_EVIDENCE_QUOTE",
    "SUPPORTED_CLAIM_QUOTE",
    "SUPPORTED_EVIDENCE_QUOTE",
    "sample_history_v1",
    "sample_history_v2",
    "sample_same_heading_fixed_body",
    "sample_new_draft",
    "sample_overclaim_draft",
    "sample_supported_claim_draft",
    "sample_long_section_draft",
    "sample_nested_coverage_draft",
    "sample_full_thesis_draft",
    "FULL_ABSTRACT_RATE",
    "FULL_CONCLUSION_RATE",
    "FULL_OVERCLAIM",
    "FULL_STATS",
]


def sample_history_v1() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行分类。")
    paragraph = doc.add_paragraph()
    paragraph.add_run("本研究")
    paragraph.add_run("非常").bold = True
    paragraph.add_run("非常")
    paragraph.add_run("有效。")
    doc.add_comment(paragraph.runs, text="避免主观评价，请给出实验依据。", author="老师甲", initials="甲")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "示例模型"
    table.cell(1, 1).text = "示例值"
    return _save(doc)


def sample_history_v2() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行图像分类。")
    paragraph = doc.add_paragraph()
    paragraph.add_run("本研究非常非常有效。")
    doc.add_comment(paragraph.runs, text="此句仍缺实验依据。", author="老师甲", initials="甲")
    return _save(doc)


def sample_same_heading_fixed_body(*, history: bool) -> bytes:
    doc = Document()
    heading = doc.add_paragraph("3.1 实验设计")
    if history:
        doc.add_comment(heading.runs, text="此处缺少实验步骤与评价指标。", author="老师甲", initials="甲")
        doc.add_paragraph("本节尚未展开。")
    else:
        doc.add_paragraph("本节给出数据集划分、训练轮次与准确率、F1 两项评价指标。")
    return _save(doc)


def sample_nested_coverage_draft() -> bytes:
    """Numbered wrappers and figure/table captions that should not look like skipped chapters."""
    doc = Document()
    doc.add_paragraph("大连财经学院本科毕业论文")
    doc.add_paragraph("摘要")
    doc.add_paragraph("摘要正文，说明研究问题与主要结果。")
    doc.add_paragraph("1 绪论")
    doc.add_paragraph("1.1 研究背景")
    doc.add_paragraph("研究背景正文，说明问题来源与研究动机。")
    doc.add_paragraph("1.2 方法概述")
    doc.add_paragraph("方法概述正文，为后文实验提供铺垫。")
    doc.add_paragraph("图 1 示意图")
    doc.add_paragraph("2 方法")
    doc.add_paragraph("2.1 模型设计")
    doc.add_paragraph("模型设计正文，给出网络结构与训练设置。")
    doc.add_paragraph("表 1 主要结果")
    doc.add_paragraph("3 结论")
    doc.add_paragraph("结论正文，回扣摘要中的主要数字。")
    return _save(doc)


def sample_long_section_draft() -> bytes:
    doc = Document()
    doc.add_paragraph("1 实验")
    for index in range(20):
        doc.add_paragraph(f"实验步骤说明段落{index:02d}，补充若干占位文字以便超过单次读取上限。")
    doc.add_paragraph("2 结论")
    doc.add_paragraph("结论段落仅作导航边界。")
    return _save(doc)


def sample_new_draft() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行图像分类。")
    doc.add_paragraph("本研究非常非常有效。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "示例模型"
    table.cell(1, 1).text = "示例值"
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 张三. 示例文献. 期刊, 2024.")
    doc.add_paragraph("[3] 李四. 另一篇文献. 期刊, 2025.")
    return _save(doc)


def _save(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
